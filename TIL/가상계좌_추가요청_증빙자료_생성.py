import mysql.connector
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

# DB 접속 정보
DB_CONFIG = {
    "host": "localhost",
    "port": 3306,
    "user": "admin",
    "password": "vraccount#$",
    "database": "vraccount",
    "charset": "utf8mb4",
    "auth_plugin": "mysql_native_password",
}

QUERY_SUMMARY = """
SELECT
    CASE
        WHEN ACCT_TYPE IS NULL THEN '미할당(NULL)'
        WHEN ACCT_TYPE = ''   THEN '미할당(빈값)'
        WHEN ACCT_TYPE = 'R'  THEN '라이더'
        WHEN ACCT_TYPE = 'S'  THEN '협력사'
        ELSE ACCT_TYPE
    END AS 계좌유형,
    CASE
        WHEN ACCT_ST = '1' THEN '정상'
        WHEN ACCT_ST = '2' THEN '정지'
        WHEN ACCT_ST = '3' THEN '해지'
        ELSE ACCT_ST
    END AS 계좌상태,
    COUNT(*) AS 건수
FROM vr_virtual_account_list
GROUP BY ACCT_TYPE, ACCT_ST
ORDER BY ACCT_TYPE, ACCT_ST
"""

QUERY_NEEDED = """
SELECT
    'RIDER' AS 구분,
    ri.RIDER_ID AS ID,
    ri.RIDER_NAME AS 이름,
    ri.SUPP_ID AS 협력사ID,
    ri.RIDER_INS_DATE AS 가입일
FROM vr_rider_info ri
LEFT JOIN vr_virtual_account_list vval
    ON ri.RIDER_ID = vval.RIDER_ID
    AND vval.ACCT_TYPE = 'R'
    AND vval.ACCT_ST = '1'
WHERE vval.RIDER_ID IS NULL
  AND ri.CONT_STATUS = 'UNDER_CONTRACT'

UNION ALL

SELECT
    'SUPPLIER' AS 구분,
    si.SUPP_ID AS ID,
    si.SUPP_NAME AS 이름,
    si.SUPP_ID AS 협력사ID,
    si.INS_DATE AS 가입일
FROM vr_supplier_info si
LEFT JOIN vr_virtual_account_list vval
    ON si.SUPP_ID = vval.SUPP_ID
    AND vval.ACCT_TYPE = 'S'
    AND vval.ACCT_ST = '1'
WHERE vval.SUPP_ID IS NULL
  AND si.STATUS_CODE = 'OPERATING'

ORDER BY 구분, 가입일
"""


def make_header_style():
    fill = PatternFill("solid", fgColor="2F5496")
    font = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    return fill, font, align, border


def make_cell_style(row_idx):
    fill = PatternFill("solid", fgColor="DCE6F1" if row_idx % 2 == 0 else "FFFFFF")
    font = Font(size=10)
    align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    return fill, font, align, border


def write_sheet(ws, title, headers, rows):
    fill_h, font_h, align_h, border_h = make_header_style()

    # 시트 제목
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=13, color="1F3864")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # 생성일시
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(row=2, column=1, value=f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ws.cell(row=2, column=1).font = Font(size=9, color="666666")
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="right")

    # 헤더
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.fill = fill_h
        c.font = font_h
        c.alignment = align_h
        c.border = border_h
    ws.row_dimensions[3].height = 22

    # 데이터
    for row_idx, row in enumerate(rows, 1):
        fill_c, font_c, align_c, border_c = make_cell_style(row_idx)
        for col, val in enumerate(row, 1):
            c = ws.cell(row=row_idx + 3, column=col, value=val)
            c.fill = fill_c
            c.font = font_c
            c.alignment = align_c
            c.border = border_c
        ws.row_dimensions[row_idx + 3].height = 20

    # 열 너비 자동 조정
    for col in range(1, len(headers) + 1):
        max_len = len(str(headers[col - 1]))
        for row in ws.iter_rows(min_row=3, min_col=col, max_col=col):
            for c in row:
                if c.value:
                    max_len = max(max_len, len(str(c.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 4, 40)


def main():
    conn = mysql.connector.connect(**DB_CONFIG)
    try:
        cur = conn.cursor()
        # 쿼리 1
        cur.execute(QUERY_SUMMARY)
        summary_rows = cur.fetchall()

        # 쿼리 2
        cur.execute(QUERY_NEEDED)
        needed_rows = cur.fetchall()
        cur.close()
    finally:
        conn.close()

    wb = openpyxl.Workbook()

    # 시트 1 - 현황 요약
    ws1 = wb.active
    ws1.title = "1. 가상계좌 현황"
    write_sheet(
        ws1,
        "가상계좌 현황 요약",
        ["계좌유형", "계좌상태", "건수"],
        summary_rows,
    )

    # 합계 행
    total_row = len(summary_rows) + 4
    total = sum(r[2] for r in summary_rows)
    c = ws1.cell(row=total_row, column=2, value="합계")
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")
    c = ws1.cell(row=total_row, column=3, value=total)
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")

    # 시트 2 - 필요 유저 목록
    ws2 = wb.create_sheet("2. 가상계좌 필요 목록")
    write_sheet(
        ws2,
        f"가상계좌 필요 유저 목록 (총 {len(needed_rows)}건)",
        ["구분", "ID", "이름", "협력사ID", "가입일"],
        needed_rows,
    )

    # 저장
    output_dir = os.path.dirname(os.path.abspath(__file__))
    filename = f"가상계좌_추가요청_증빙_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output_path = os.path.join(output_dir, filename)
    wb.save(output_path)
    print(f"저장 완료: {output_path}")


if __name__ == "__main__":
    main()
